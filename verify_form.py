#!/usr/bin/env python3
import re, sys
import openpyxl
from openpyxl.utils import get_column_letter

FN=re.compile(r"([A-Za-z_][A-Za-z0-9_.]*)\s*\(")
WHITELIST={"IF","SUM","COUNT","COUNTA","COUNTIF","COUNTIFS","SUMPRODUCT","DATE",
           "INDEX","MATCH","IFERROR","MID","LEFT","RIGHT","SUBSTITUTE","TRIM","LEN"}
wb=openpyxl.load_workbook("IATSE_479_Field_Visit_Form.xlsx")
log=wb["Visit Log"]
headers={get_column_letter(c):log.cell(row=1,column=c).value for c in range(1,log.max_column+1)}
def col(h):
    for L,hh in headers.items():
        if hh==h: return L
    raise KeyError(h)

errors=[]; n=0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            v=cell.value
            if not (isinstance(v,str) and v.startswith("=")): continue
            n+=1; loc=f"{ws.title}!{cell.coordinate}"
            if v.count("(")!=v.count(")"): errors.append(f"{loc}: unbalanced parens")
            if v.count('"')%2: errors.append(f"{loc}: odd number of quotes")
            for fn in FN.findall(v):
                base=fn.split(".")[-1].upper()
                if base not in WHITELIST: errors.append(f"{loc}: non-whitelisted fn {fn}")
            for sheetref in ["Visit Log","Field Rep Summary"]:
                if sheetref in v and f"'{sheetref}'" not in v:
                    errors.append(f"{loc}: unquoted '{sheetref}' ref")
print("Formulas scanned:",n)
print("Structural errors:",len(errors))
for e in errors[:40]: print("  ",e)

# example row recompute
def val(h): return log[f"{col(h)}2"].value
DEPTS=["Art","Construction","Costumes","Craft Service","Electrical","Greens","Grip","Locations",
       "Medic","Paint","Plaster","Props","Set Dec","Set Teacher","Sound","SPFX","Video"]
depts_seen=", ".join(d for d in DEPTS if str(val(f"Dept: {d}")).strip().lower()=="yes")
checks=[
 ("Production",val("Production Title"),"Northern Lights (S2)"),
 ("Location Visited",val("Location Visited"),"Stage 4, Pinewood Studios"),
 ("Proof of Visit",val("Proof of Visit"),"Emailed production"),
 ("Departments Seen (derived)",depts_seen,"Electrical, Grip, Set Dec"),
]
print("\nField Rep Summary expected values for the example visit:")
ok=True
for name,got,exp in checks:
    good=str(got)==exp; ok&=good
    print(f"  [{'OK' if good else 'XX'}] {name}: {got!r}  (expected {exp!r})")

# check Visit Form references resolve to real Visit Log columns
vf=wb["Visit Form"]
missing=[]
for row in vf.iter_rows():
    for cell in row:
        v=cell.value
        if isinstance(v,str) and v.startswith("="):
            for m in re.findall(r"'Visit Log'!\$([A-Z]{1,3})\$2", v):
                if m not in headers or headers[m] is None: missing.append((cell.coordinate,m))
print("\nVisit Form -> Visit Log column refs that are out of range:", len(missing))
for c,m in missing[:10]: print("  ",c,m)

print("\nRESULT:", "PASS" if (not errors and ok and not missing) else "FAIL")
sys.exit(0 if (not errors and ok and not missing) else 1)
