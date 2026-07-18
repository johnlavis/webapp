#!/usr/bin/env python3
"""
Import filled IATSE Local 479 X-4 Field Visit PDFs into an Excel work-activity database.

Usage:
    python import_visits.py  visit1.pdf visit2.pdf ...
    python import_visits.py  "C:/path/to/FilledForms/*.pdf"
    python import_visits.py                 # imports every *.pdf in ./FilledForms

Each filled PDF becomes one row in Field_Rep_Database.xlsx (created if missing).
Re-importing the same file name is skipped, so you can point it at a growing folder.
"""
import sys, glob, os, datetime
from pypdf import PdfReader
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

DB = "Field_Rep_Database.xlsx"
NAVY = "1F3864"; GREEN = "1E7145"; WHITE = "FFFFFF"; STRIPE = "EAF3EA"
thin = Side(style="thin", color="B0B0B0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# Database columns (order matters). The 5 you asked to track are first.
COLUMNS = ["Date of Visit", "Production", "Location Visited", "Departments Seen",
           "Proof of Visit", "Field Representative", "Production Phase",
           "Reason for Visit", "Arrival", "Departure", "Source PDF", "Imported"]

AREA_FIELDS = {"Visited": "Stage", "Location": "Location", "Mill": "Mill",
               "Lockup": "Lockup", "Other Location": "Other"}

def txt(fields, key):
    v = fields.get(key, {}).get("/V")
    return "" if v is None else str(v).strip()

def on(fields, key):
    """Return the checkbox's on-label if checked, else None."""
    v = fields.get(key, {}).get("/V")
    if v is None: return None
    s = str(v)
    return None if s in ("/Off", "Off", "") else s.lstrip("/")

def extract(pdf_path):
    fields = PdfReader(pdf_path).get_fields() or {}
    # locations = the Areas Visited checkboxes that are ticked
    locs = [label for fld, label in AREA_FIELDS.items() if on(fields, fld)]
    # departments = every DeptVisit* checkbox that is ticked (uses the tick's own label)
    depts = [on(fields, k) for k in fields if k.startswith("DeptVisit") and on(fields, k)]
    # proof: the X-4 has no "emailed" field, so default it (edit later in Excel if needed)
    proof = "Emailed production"
    return {
        "Date of Visit": txt(fields, "Date of Visit"),
        "Production": txt(fields, "Production Title"),
        "Location Visited": ", ".join(locs),
        "Departments Seen": ", ".join(sorted(depts)),
        "Proof of Visit": proof,
        "Field Representative": txt(fields, "Field Representative"),
        "Production Phase": on(fields, "ProductionPhase") or "",
        "Reason for Visit": txt(fields, "Reason for Visit routine  issue"),
        "Arrival": txt(fields, "Arrival Time"),
        "Departure": txt(fields, "Departure Time"),
        "Source PDF": os.path.basename(pdf_path),
        "Imported": datetime.date.today().isoformat(),
    }

def load_or_create():
    if os.path.exists(DB):
        wb = openpyxl.load_workbook(DB)
        return wb, wb["Visits"]
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Visits"
    ws.sheet_view.showGridLines = False
    widths = [13, 26, 22, 34, 20, 18, 14, 18, 11, 11, 30, 12]
    for i, (h, w) in enumerate(zip(COLUMNS, widths), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name="Arial", bold=True, color=WHITE, size=10)
        c.fill = PatternFill("solid", fgColor=(GREEN if i <= 5 else NAVY))
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"
    return wb, ws

def existing_sources(ws):
    src_col = COLUMNS.index("Source PDF") + 1
    return {ws.cell(row=r, column=src_col).value
            for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=src_col).value}

def main(argv):
    args = argv[1:]
    if not args:
        args = sorted(glob.glob("FilledForms/*.pdf"))
    paths = []
    for a in args:
        paths += sorted(glob.glob(a)) if any(ch in a for ch in "*?[") else [a]
    paths = [p for p in paths if p.lower().endswith(".pdf") and os.path.exists(p)]
    if not paths:
        print("No PDF files found. Pass file paths or put PDFs in a ./FilledForms folder."); return 1

    wb, ws = load_or_create()
    seen = existing_sources(ws)
    added = skipped = 0
    for p in paths:
        base = os.path.basename(p)
        if base in seen:
            print(f"  skip (already imported): {base}"); skipped += 1; continue
        row = extract(p)
        r = ws.max_row + 1
        for i, h in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=i, value=row[h])
            cell.font = Font(name="Arial", size=10)
            cell.border = border
            if r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=STRIPE)
        seen.add(base); added += 1
        print(f"  imported: {base}  ->  {row['Date of Visit'] or '(no date)'} | {row['Production'] or '(no production)'}")

    # (re)build the table over the used range
    if "VisitDB" in ws.tables: del ws.tables["VisitDB"]
    last = get_column_letter(len(COLUMNS))
    tbl = Table(displayName="VisitDB", ref=f"A1:{last}{max(ws.max_row,2)}")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tbl)
    wb.save(DB)
    print(f"\nDone. Added {added}, skipped {skipped}. Database: {DB} (now {ws.max_row-1} visit rows).")
    return 0

if __name__ == "__main__":
    sys.exit(main(sys.argv))
